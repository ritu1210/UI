"""Load and expose reference data, from Azure SQL with a CSV fallback.

Backend is chosen by config.DATA_BACKEND:
  * "auto" (default) tries SQL and falls back to the CSV dump files if the
    database is unreachable (and stays on CSV for the rest of the process);
  * "sql" forces the database;
  * "csv" forces the CSV files and never touches the database.

Reads are cached for the process lifetime; write helpers clear the caches.
The column maps below are the single source of truth shared with the one-time
migration script (scripts/migrate_to_db.py) and map each database column to its
original CSV header.
"""
from __future__ import annotations

import csv
import re
from dataclasses import dataclass
from functools import lru_cache

from . import config, db


# ---------------------------------------------------------------------------
# Domain models (unchanged public shape).
# ---------------------------------------------------------------------------
@dataclass
class Employee:
    name: str
    director: str
    job_title: str
    job_grade: str
    country: str
    location: str
    gender: str
    status: str
    start_date: str
    reporting_manager: str
    employment_type: str
    email: str


@dataclass
class Project:
    project_id: str
    director: str
    spoc: str
    program_manager: str
    title: str
    project_type: str
    commodity: str
    bu: str
    cluster: str
    current_il: str
    is_active: str = ""
    funnel_total: float = 0.0
    actual_total: float = 0.0


@dataclass
class FteAllocation:
    employee: str
    project_id: str
    project_title: str
    month: str
    allocation: float


# ---------------------------------------------------------------------------
# Column maps: database column -> original CSV header.
# The write paths receive dicts keyed by these CSV headers (from routers), and
# the migration reads the CSVs by these headers too.
# ---------------------------------------------------------------------------
PROJECT_COLUMNS: dict[str, str] = {
    "project_id": "SMRS / Project ID",
    "title": "STET Funnel Project Title",
    "director": "Director",
    "spoc": "STET SPOC",
    "program_manager": "Program Manager",
    "project_type": "Project Type (PRODUCTIVITY, AOS, LCM, NPI, CART PDC, TEST ENG, CONQ, IGM)",
    "procurement_type": "Procurement Type (PROCUREMENT, TCO, N/A)",
    "savings_type": "Savings Type (TCO, CONCEPT, SOURCING, NEGO, NON_12NC, NPP, CONQ, FCP & PPV)",
    "commodity": "Commodity",
    "bu": "BU",
    "cluster": "Cluster",
    "current_il": "Current IL",
    "il5_date": "IL5 Date",
    "is_active": "Is Active",
    "parts_dual_sourced_passive": "%23 of Parts Dual Sourced (PASSIVE)",
    "parts_dual_sourced_active": "%23 of Parts Dual Sourced (ACTIVE)",
    "aos_impact": "AOS Impact  (Enter €0 if not AOS Project)",
    "aos_impact_approved": "AOS Impact Approved or Not",
    "qn_reduction": "QN Reduction Impact (Enter €0 if not CONQ Project)",
    "impacted_parts": "Impacted Parts Added? (YES_NO_N/A) Needs to be YES or N/A for all IL5 Projects",
    "sqe_resources": "Are Any STET SQE Resources Applied to Project? (YES_NO_N/A)",
    "comments": "Comments / Challenges",
    "week": "Week",
    "funnel_2022": "2022 Funnel (Euro)",
    "funnel_2023": "2023 Funnel (Euro)",
    "funnel_2024": "2024 Funnel (Euro)",
    "funnel_2025": "2025 Funnel (Euro)",
    "funnel_2026": "2026 Funnel (Euro)",
    "funnel_2027": "2027 Funnel (Euro)",
    "funnel_2028": "2028 Funnel (Euro)",
    "actual_2019": "2019 Actual (Euro)",
    "actual_2020": "2020 Actual (Euro)",
    "actual_2021": "2021 Actual (Euro)",
    "actual_2022": "2022 Actual (Euro)",
    "actual_2023": "2023 Actual (Euro)",
    "actual_2024": "2024 Actual (Euro)",
    "actual_2025": "2025 Actual (Euro)",
    "actual_2026": "2026 Actual (Euro)",
    "actual_2027": "2027 Actual (Euro)",
    "actual_2028": "2028 Actual (Euro)",
    "source_id": "ID",
    "created": "Created",
    "created_by": "Created By",
    "modified": "Modified",
    "modified_by": "Modified By",
    "supplier_transfer": "Supplier Transfer",
    "transfer_reason": "Transfer reason",
}

EMPLOYEE_COLUMNS: dict[str, str] = {
    "name": "Employee Name (HC)",
    "email": "Email ID",
    "director": "Director",
    "job_title": "Job Title (HC)",
    "job_grade": "Job Grade (HC)",
    "country": "Location Country (HC)",
    "location": "Job Location (HC)",
    "gender": "Diversity (HC)",
    "status": "Employment Status (HC)",
    "start_date": "Employee Start Date (M/D/YYYY)",
    "reporting_manager": "Reporting Manager (HC)",
    "employment_type": "Employment Type (HC)",
    "requisition": "Requisition (HC)",
    "contractor": "Contractor (HC)",
    "comments": "Comments (HC)",
    "created": "Created",
    "created_by": "Created By",
    "modified_by": "Modified By",
}

FTE_COLUMNS: dict[str, str] = {
    "project_id": "FTE Project ID",
    "employee_name": "FTE Employee Name",
    "project_title": "FTE Project ID:FTE Project Title (Text)",
    "month_year": "MonthYearColumn",
    "allocation": "FTE Allocation Value (5% - 100%)",
    "start_date": "FTE Allocation Start Date (MM/DD/YYYY)",
    "end_date": "FTE Allocation End Date (MM/DD/YYYY)",
    "commodity": "FTE Project ID:FTE Commodity (Text)",
    "spoc": "FTE Project ID:FTE STET SPOC (Text)",
    "procurement_type": "FTE Project ID:FTE Procurement Type (Text)",
    "current_il": "FTE Project ID:FTE Current IL (Text)",
    "savings_type": "FTE Project ID:FTE Savings Type (Text)",
    "director": "FTE Project ID:FTE Director (Text)",
    "project_type": "FTE Project ID:FTE Project Type (Text)",
    "bu": "FTE Project ID:FTE BU (Text)",
    "cluster": "FTE Project ID:FTE Cluster (Text)",
    "reporting_manager": "FTE Employee Name: Reporting Manager (HC)",
    "approval_status": "Approval Status",
    "source_id": "ID",
    "created": "Created",
    "created_by": "Created By",
    "modified": "Modified",
    "modified_by": "Modified By",
}

# Euro/decimal columns are stored numeric and parsed with _euro on write.
PROJECT_NUMERIC_COLUMNS = {
    c for c in PROJECT_COLUMNS if c.startswith(("funnel_", "actual_"))
}

# Reverse maps (CSV header -> database column) for the write paths.
_PROJECT_CSV_TO_DB = {csv_h: dbcol for dbcol, csv_h in PROJECT_COLUMNS.items()}
_EMPLOYEE_CSV_TO_DB = {csv_h: dbcol for dbcol, csv_h in EMPLOYEE_COLUMNS.items()}

# Funnel/actual columns that make up the headline totals shown in the UI.
_FUNNEL_TOTAL_COLS = [f"funnel_{y}" for y in range(2025, 2029)]
_ACTUAL_TOTAL_COLS = [f"actual_{y}" for y in range(2025, 2029)]

# CSV headers for the totals (used by the CSV backend).
_FUNNEL_EURO_HEADERS = [PROJECT_COLUMNS[c] for c in _FUNNEL_TOTAL_COLS]
_ACTUAL_EURO_HEADERS = [PROJECT_COLUMNS[c] for c in _ACTUAL_TOTAL_COLS]


# ---------------------------------------------------------------------------
# Backend selection (SQL with automatic CSV fallback).
# ---------------------------------------------------------------------------
_ACTIVE_BACKEND: str | None = None


def active_backend() -> str:
    """Return the backend in use ("sql" or "csv"), deciding once per process."""
    global _ACTIVE_BACKEND
    if _ACTIVE_BACKEND is None:
        if config.DATA_BACKEND in ("sql", "csv"):
            _ACTIVE_BACKEND = config.DATA_BACKEND
        else:  # "auto": probe the database, fall back to CSV on failure
            _ACTIVE_BACKEND = _probe_sql()
    return _ACTIVE_BACKEND


def _probe_sql() -> str:
    try:
        conn = db.get_connection()
        conn.close()
        return "sql"
    except Exception:
        return "csv"


def _use_csv() -> None:
    """Latch the process onto the CSV backend after a database failure."""
    global _ACTIVE_BACKEND
    _ACTIVE_BACKEND = "csv"


def reset_backend() -> None:
    """Forget the cached backend choice and clear data caches (for testing)."""
    global _ACTIVE_BACKEND
    _ACTIVE_BACKEND = None
    for fn in (load_employees, load_projects, project_index, employee_index, load_fte_allocations):
        fn.cache_clear()


# ---------------------------------------------------------------------------
# Value helpers.
# ---------------------------------------------------------------------------
def _clean(value) -> str:
    """Normalise a raw cell (strip stray whitespace / newlines)."""
    if value is None:
        return ""
    return " ".join(str(value).split()).strip()


def _euro(value) -> float:
    """Parse a euro cell like '€ 314,535.00' into a float (0.0 if empty/invalid)."""
    if value is None or value == "":
        return 0.0
    s = "".join(ch for ch in str(value) if ch.isdigit() or ch in ".-")
    if s in ("", "-", ".", "-."):
        return 0.0
    try:
        return float(s)
    except ValueError:
        return 0.0


def _parse_pct(value) -> float | None:
    txt = _clean(value).replace("%", "")
    if not txt:
        return None
    try:
        return float(txt)
    except ValueError:
        return None


def _num_str(value) -> str:
    """Render a numeric DB value back to a plain string for edit forms."""
    if value is None:
        return ""
    return _clean(value)


# ---------------------------------------------------------------------------
# Public reads (dispatch SQL -> CSV fallback; cached).
# ---------------------------------------------------------------------------
@lru_cache(maxsize=1)
def load_employees() -> list[Employee]:
    if active_backend() == "csv":
        return _csv_load_employees()
    try:
        return _sql_load_employees()
    except Exception:
        _use_csv()
        return _csv_load_employees()


@lru_cache(maxsize=1)
def load_projects() -> list[Project]:
    if active_backend() == "csv":
        return _csv_load_projects()
    try:
        return _sql_load_projects()
    except Exception:
        _use_csv()
        return _csv_load_projects()


@lru_cache(maxsize=1)
def load_fte_allocations() -> list[FteAllocation]:
    if active_backend() == "csv":
        return _csv_load_fte()
    try:
        return _sql_load_fte()
    except Exception:
        _use_csv()
        return _csv_load_fte()


@lru_cache(maxsize=1)
def project_index() -> dict[str, Project]:
    return {p.project_id: p for p in load_projects()}


@lru_cache(maxsize=1)
def employee_index() -> dict[str, Employee]:
    return {e.name: e for e in load_employees()}


def business_units() -> list[str]:
    units = {p.bu for p in load_projects() if p.bu}
    return sorted(units)


# ---------------------------------------------------------------------------
# KPI dashboard records (Financials + Projects dashboards). Each record carries
# the descriptive fields plus per-year funnel/actual euro values so the frontend
# can slice and aggregate exactly like the Power BI report.
# ---------------------------------------------------------------------------
KPI_FUNNEL_YEARS = list(range(2022, 2029))
KPI_ACTUAL_YEARS = list(range(2019, 2029))
_KPI_DESCRIPTIVE = [
    "project_id", "title", "director", "spoc", "program_manager", "bu",
    "cluster", "commodity", "project_type", "procurement_type", "savings_type",
    "current_il", "is_active", "il5_date",
]


def _kpi_build(get) -> dict:
    """Build one KPI record; `get(db_col)` returns the raw value for a column."""
    m = re.search(r"(20\d{2})", _clean(get("il5_date")))
    funnel = {str(y): _euro(get(f"funnel_{y}")) for y in KPI_FUNNEL_YEARS}
    actual = {str(y): _euro(get(f"actual_{y}")) for y in KPI_ACTUAL_YEARS}
    return {
        "project_id": _clean(get("project_id")),
        "title": _clean(get("title")),
        "director": _clean(get("director")),
        "spoc": _clean(get("spoc")),
        "program_manager": _clean(get("program_manager")),
        "bu": _clean(get("bu")),
        "cluster": _clean(get("cluster")),
        "commodity": _clean(get("commodity")),
        "project_type": _clean(get("project_type")),
        "procurement_type": _clean(get("procurement_type")),
        "savings_type": _clean(get("savings_type")),
        "current_il": _clean(get("current_il")),
        "is_active": _clean(get("is_active")) or "Yes",
        "il5_year": int(m.group(1)) if m else None,
        "funnel": funnel,
        "actual": actual,
        "funnel_total": sum(funnel[str(y)] for y in range(2025, 2029)),
        "actual_total": sum(actual[str(y)] for y in range(2025, 2029)),
    }


def kpi_records() -> list[dict]:
    if active_backend() == "csv":
        return _csv_kpi_records()
    try:
        return _sql_kpi_records()
    except Exception:
        _use_csv()
        return _csv_kpi_records()


def _sql_kpi_records() -> list[dict]:
    cols = (_KPI_DESCRIPTIVE
            + [f"funnel_{y}" for y in KPI_FUNNEL_YEARS]
            + [f"actual_{y}" for y in KPI_ACTUAL_YEARS])
    sql = f"SELECT {', '.join(f'[{c}]' for c in cols)} FROM [{config.PROJECTS_TABLE}]"
    conn = db.get_connection()
    try:
        cur = conn.cursor()
        cur.execute(sql)
        rows = cur.fetchall()
    finally:
        conn.close()
    out: list[dict] = []
    seen: set[str] = set()
    for row in rows:
        rec = dict(zip(cols, row))
        pid = _clean(rec.get("project_id"))
        if not pid or pid in seen:
            continue
        seen.add(pid)
        out.append(_kpi_build(rec.get))
    return out


def _csv_kpi_records() -> list[dict]:
    out: list[dict] = []
    seen: set[str] = set()
    with config.FUNNEL_FILE.open(encoding="utf-8-sig", newline="") as fh:
        for row in csv.DictReader(fh):
            def get(col, _row=row):
                return _row.get(PROJECT_COLUMNS[col])
            pid = _clean(get("project_id"))
            if not pid or pid in seen:
                continue
            seen.add(pid)
            out.append(_kpi_build(get))
    return out


# ---------------------------------------------------------------------------
# Resource-allocation records (FTE export workbook) for the Resources dashboard.
# ---------------------------------------------------------------------------
_RESOURCE_MAP = {
    "project_id": "ProjectID",
    "poc": "STET POC",
    "project_name": "ProjectName",
    "director": "Directors",
    "procurement_type": "ProcurementType",
    "project_type": "ProjectTypeName",
    "savings_type": "SavingsCategoryName",
    "commodity": "CommodityName",
    "bu": "BuName",
    "cluster": "ClusterName",
    "phase": "PhaseName",
    "person": "FTE",
    "fte_type": "FTE/Contingent",
    "region": "Region",
    "reporting_manager": "FTE: Reporting Manager",
    "fte_director": "FTE:Director",
    "proj_others": "Projects/Others",
    "allocation_status": "Allocation Status",
    "status": "Status",
}


@lru_cache(maxsize=1)
def resource_records() -> list[dict]:
    """Read the FTE export workbook into resource-allocation records."""
    from openpyxl import load_workbook

    wb = load_workbook(config.RESOURCE_FILE, read_only=True, data_only=True)
    try:
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else "" for h in next(rows)]
        idx = {name: i for i, name in enumerate(header)}

        def cell(row, header_name):
            i = idx.get(header_name)
            return row[i] if i is not None and i < len(row) else None

        out: list[dict] = []
        for row in rows:
            rec = {key: _clean(cell(row, col)) for key, col in _RESOURCE_MAP.items()}
            rec["resource_amount"] = _euro(cell(row, "ResourceAmount"))
            if not (rec["project_name"] or rec["person"] or rec["project_id"]):
                continue
            out.append(rec)
    finally:
        wb.close()
    return out


def _clear_project_caches() -> None:
    load_projects.cache_clear()
    project_index.cache_clear()


def _clear_employee_caches() -> None:
    load_employees.cache_clear()
    employee_index.cache_clear()


# ---------------------------------------------------------------------------
# SQL backend — reads.
# ---------------------------------------------------------------------------
def _sql_load_employees() -> list[Employee]:
    cols = ["name", "director", "job_title", "job_grade", "country", "location",
            "gender", "status", "start_date", "reporting_manager",
            "employment_type", "email"]
    sql = f"SELECT {', '.join(f'[{c}]' for c in cols)} FROM [{config.EMPLOYEES_TABLE}]"
    conn = db.get_connection()
    try:
        cur = conn.cursor()
        cur.execute(sql)
        rows = cur.fetchall()
    finally:
        conn.close()
    employees = [
        Employee(**{c: _clean(v) for c, v in zip(cols, row)})
        for row in rows
        if _clean(row[0])
    ]
    employees.sort(key=lambda e: e.name.lower())
    return employees


def _sql_load_projects() -> list[Project]:
    base = ["project_id", "director", "spoc", "program_manager", "title",
            "project_type", "commodity", "bu", "cluster", "current_il", "is_active"]
    cols = base + _FUNNEL_TOTAL_COLS + _ACTUAL_TOTAL_COLS
    sql = f"SELECT {', '.join(f'[{c}]' for c in cols)} FROM [{config.PROJECTS_TABLE}]"
    conn = db.get_connection()
    try:
        cur = conn.cursor()
        cur.execute(sql)
        rows = cur.fetchall()
    finally:
        conn.close()

    projects: list[Project] = []
    for row in rows:
        rec = dict(zip(cols, row))
        pid = _clean(rec["project_id"])
        if not pid:
            continue
        projects.append(
            Project(
                project_id=pid,
                director=_clean(rec["director"]),
                spoc=_clean(rec["spoc"]),
                program_manager=_clean(rec["program_manager"]),
                title=_clean(rec["title"]),
                project_type=_clean(rec["project_type"]),
                commodity=_clean(rec["commodity"]),
                bu=_clean(rec["bu"]),
                cluster=_clean(rec["cluster"]),
                current_il=_clean(rec["current_il"]),
                is_active=_clean(rec["is_active"]),
                funnel_total=sum(_euro(rec[c]) for c in _FUNNEL_TOTAL_COLS),
                actual_total=sum(_euro(rec[c]) for c in _ACTUAL_TOTAL_COLS),
            )
        )
    projects.sort(key=lambda p: p.project_id.lower())
    return projects


def _sql_load_fte() -> list[FteAllocation]:
    sql = (
        f"SELECT [employee_name], [project_id], [project_title], [month_year], "
        f"[allocation] FROM [{config.FTE_ALLOCATIONS_TABLE}]"
    )
    conn = db.get_connection()
    try:
        cur = conn.cursor()
        cur.execute(sql)
        rows = cur.fetchall()
    finally:
        conn.close()

    allocs: list[FteAllocation] = []
    for employee, project_id, title, month, pct in rows:
        employee = _clean(employee)
        project_id = _clean(project_id)
        month = _clean(month)
        if not (employee and project_id and month) or pct is None:
            continue
        allocs.append(
            FteAllocation(
                employee=employee,
                project_id=project_id,
                project_title=_clean(title),
                month=month,
                allocation=float(pct),
            )
        )
    return allocs


# ---------------------------------------------------------------------------
# CSV backend — reads.
# ---------------------------------------------------------------------------
def _csv_load_employees() -> list[Employee]:
    employees: list[Employee] = []
    with config.HEADCOUNT_FILE.open(encoding="utf-8-sig", newline="") as fh:
        for row in csv.DictReader(fh):
            name = _clean(row.get(EMPLOYEE_COLUMNS["name"]))
            if not name:
                continue
            employees.append(
                Employee(
                    name=name,
                    director=_clean(row.get(EMPLOYEE_COLUMNS["director"])),
                    job_title=_clean(row.get(EMPLOYEE_COLUMNS["job_title"])),
                    job_grade=_clean(row.get(EMPLOYEE_COLUMNS["job_grade"])),
                    country=_clean(row.get(EMPLOYEE_COLUMNS["country"])),
                    location=_clean(row.get(EMPLOYEE_COLUMNS["location"])),
                    gender=_clean(row.get(EMPLOYEE_COLUMNS["gender"])),
                    status=_clean(row.get(EMPLOYEE_COLUMNS["status"])),
                    start_date=_clean(row.get(EMPLOYEE_COLUMNS["start_date"])),
                    reporting_manager=_clean(row.get(EMPLOYEE_COLUMNS["reporting_manager"])),
                    employment_type=_clean(row.get(EMPLOYEE_COLUMNS["employment_type"])),
                    email=_clean(row.get(EMPLOYEE_COLUMNS["email"])),
                )
            )
    employees.sort(key=lambda e: e.name.lower())
    return employees


def _csv_load_projects() -> list[Project]:
    projects: list[Project] = []
    seen: set[str] = set()
    with config.FUNNEL_FILE.open(encoding="utf-8-sig", newline="") as fh:
        for row in csv.DictReader(fh):
            pid = _clean(row.get(PROJECT_COLUMNS["project_id"]))
            if not pid or pid in seen:
                continue
            seen.add(pid)
            projects.append(
                Project(
                    project_id=pid,
                    director=_clean(row.get(PROJECT_COLUMNS["director"])),
                    spoc=_clean(row.get(PROJECT_COLUMNS["spoc"])),
                    program_manager=_clean(row.get(PROJECT_COLUMNS["program_manager"])),
                    title=_clean(row.get(PROJECT_COLUMNS["title"])),
                    project_type=_clean(row.get(PROJECT_COLUMNS["project_type"])),
                    commodity=_clean(row.get(PROJECT_COLUMNS["commodity"])),
                    bu=_clean(row.get(PROJECT_COLUMNS["bu"])),
                    cluster=_clean(row.get(PROJECT_COLUMNS["cluster"])),
                    current_il=_clean(row.get(PROJECT_COLUMNS["current_il"])),
                    is_active=_clean(row.get(PROJECT_COLUMNS["is_active"])),
                    funnel_total=sum(_euro(row.get(c)) for c in _FUNNEL_EURO_HEADERS),
                    actual_total=sum(_euro(row.get(c)) for c in _ACTUAL_EURO_HEADERS),
                )
            )
    projects.sort(key=lambda p: p.project_id.lower())
    return projects


def _csv_load_fte() -> list[FteAllocation]:
    allocs: list[FteAllocation] = []
    with config.FTE_ALLOCATION_FILE.open(encoding="utf-8-sig", newline="") as fh:
        for row in csv.DictReader(fh):
            employee = _clean(row.get(FTE_COLUMNS["employee_name"]))
            project_id = _clean(row.get(FTE_COLUMNS["project_id"]))
            month = _clean(row.get(FTE_COLUMNS["month_year"]))
            pct = _parse_pct(row.get(FTE_COLUMNS["allocation"]))
            if not (employee and project_id and month) or pct is None:
                continue
            allocs.append(
                FteAllocation(
                    employee=employee,
                    project_id=project_id,
                    project_title=_clean(row.get(FTE_COLUMNS["project_title"])),
                    month=month,
                    allocation=pct,
                )
            )
    return allocs


# ---------------------------------------------------------------------------
# Public writes. `values`/`updates` are keyed by the original CSV headers
# (routers use FUNNEL_COLUMNS / HEADCOUNT_COLUMNS).
# ---------------------------------------------------------------------------
def append_project(values: dict[str, str]) -> Project:
    if active_backend() == "csv":
        return _csv_append_project(values)
    try:
        return _sql_append_project(values)
    except Exception:
        _use_csv()
        return _csv_append_project(values)


def append_projects(rows: list[dict[str, str]]) -> int:
    if not rows:
        return 0
    if active_backend() == "csv":
        return _csv_append_projects(rows)
    try:
        return _sql_append_projects(rows)
    except Exception:
        _use_csv()
        return _csv_append_projects(rows)


def funnel_row(project_id: str) -> dict[str, str] | None:
    if active_backend() == "csv":
        return _csv_funnel_row(project_id)
    try:
        return _sql_funnel_row(project_id)
    except Exception:
        _use_csv()
        return _csv_funnel_row(project_id)


def update_project(project_id: str, updates: dict[str, str]) -> Project | None:
    if active_backend() == "csv":
        return _csv_update_project(project_id, updates)
    try:
        return _sql_update_project(project_id, updates)
    except Exception:
        _use_csv()
        return _csv_update_project(project_id, updates)


def append_employee(values: dict[str, str]) -> Employee:
    if active_backend() == "csv":
        return _csv_append_employee(values)
    try:
        return _sql_append_employee(values)
    except Exception:
        _use_csv()
        return _csv_append_employee(values)


def employee_row(key: str) -> dict[str, str] | None:
    if active_backend() == "csv":
        return _csv_employee_row(key)
    try:
        return _sql_employee_row(key)
    except Exception:
        _use_csv()
        return _csv_employee_row(key)


def update_employee(key: str, updates: dict[str, str]) -> Employee | None:
    if active_backend() == "csv":
        return _csv_update_employee(key, updates)
    try:
        return _sql_update_employee(key, updates)
    except Exception:
        _use_csv()
        return _csv_update_employee(key, updates)


# ---------------------------------------------------------------------------
# SQL backend — writes.
# ---------------------------------------------------------------------------
def _project_db_values(values: dict[str, str]) -> dict[str, object]:
    """Translate a CSV-header-keyed dict into {db_column: value}."""
    out: dict[str, object] = {}
    for csv_header, val in values.items():
        col = _PROJECT_CSV_TO_DB.get(csv_header)
        if not col:
            continue
        out[col] = _euro(val) if col in PROJECT_NUMERIC_COLUMNS else _clean(val)
    return out


def _insert_row(conn, table: str, db_values: dict[str, object]) -> None:
    cols = list(db_values.keys())
    placeholders = ", ".join("?" for _ in cols)
    col_sql = ", ".join(f"[{c}]" for c in cols)
    conn.cursor().execute(
        f"INSERT INTO [{table}] ({col_sql}) VALUES ({placeholders})",
        [db_values[c] for c in cols],
    )


def _sql_append_project(values: dict[str, str]) -> Project:
    db_values = _project_db_values(values)
    conn = db.get_connection()
    try:
        _insert_row(conn, config.PROJECTS_TABLE, db_values)
        conn.commit()
    finally:
        conn.close()
    _clear_project_caches()
    return project_index()[db_values["project_id"]]


def _sql_append_projects(rows: list[dict[str, str]]) -> int:
    conn = db.get_connection()
    try:
        cur = conn.cursor()
        for values in rows:
            db_values = _project_db_values(values)
            if not db_values.get("project_id"):
                continue
            cols = list(db_values.keys())
            placeholders = ", ".join("?" for _ in cols)
            col_sql = ", ".join(f"[{c}]" for c in cols)
            cur.execute(
                f"INSERT INTO [{config.PROJECTS_TABLE}] ({col_sql}) VALUES ({placeholders})",
                [db_values[c] for c in cols],
            )
        conn.commit()
    finally:
        conn.close()
    _clear_project_caches()
    return len(rows)


def _sql_funnel_row(project_id: str) -> dict[str, str] | None:
    cols = list(PROJECT_COLUMNS.keys())
    sql = (
        f"SELECT {', '.join(f'[{c}]' for c in cols)} "
        f"FROM [{config.PROJECTS_TABLE}] WHERE [project_id] = ?"
    )
    conn = db.get_connection()
    try:
        cur = conn.cursor()
        cur.execute(sql, [project_id])
        row = cur.fetchone()
    finally:
        conn.close()
    if row is None:
        return None
    return {PROJECT_COLUMNS[col]: _num_str(val) for col, val in zip(cols, row)}


def _sql_update_project(project_id: str, updates: dict[str, str]) -> Project | None:
    db_values = _project_db_values(updates)
    db_values.pop("project_id", None)
    if not db_values:
        return project_index().get(project_id)
    set_sql = ", ".join(f"[{c}] = ?" for c in db_values)
    params = list(db_values.values()) + [project_id]
    conn = db.get_connection()
    try:
        cur = conn.cursor()
        cur.execute(
            f"UPDATE [{config.PROJECTS_TABLE}] SET {set_sql} WHERE [project_id] = ?",
            params,
        )
        affected = cur.rowcount
        conn.commit()
    finally:
        conn.close()
    if not affected:
        return None
    _clear_project_caches()
    return project_index().get(project_id)


def _sql_append_employee(values: dict[str, str]) -> Employee:
    db_values = {
        _EMPLOYEE_CSV_TO_DB[csv_h]: _clean(val)
        for csv_h, val in values.items()
        if csv_h in _EMPLOYEE_CSV_TO_DB
    }
    conn = db.get_connection()
    try:
        _insert_row(conn, config.EMPLOYEES_TABLE, db_values)
        conn.commit()
    finally:
        conn.close()
    _clear_employee_caches()
    return employee_index()[db_values["name"]]


def _sql_employee_row(key: str) -> dict[str, str] | None:
    cols = list(EMPLOYEE_COLUMNS.keys())
    select = ", ".join(f"[{c}]" for c in cols)
    where = "[email] = ?" if "@" in key else "[name] = ?"
    conn = db.get_connection()
    try:
        cur = conn.cursor()
        cur.execute(
            f"SELECT {select} FROM [{config.EMPLOYEES_TABLE}] WHERE {where}", [key]
        )
        row = cur.fetchone()
    finally:
        conn.close()
    if row is None:
        return None
    return {EMPLOYEE_COLUMNS[col]: _clean(val) for col, val in zip(cols, row)}


def _sql_update_employee(key: str, updates: dict[str, str]) -> Employee | None:
    db_values = {
        _EMPLOYEE_CSV_TO_DB[csv_h]: _clean(val)
        for csv_h, val in updates.items()
        if csv_h in _EMPLOYEE_CSV_TO_DB
    }
    where = "[email] = ?" if "@" in key else "[name] = ?"
    conn = db.get_connection()
    try:
        cur = conn.cursor()
        cur.execute(
            f"SELECT [name] FROM [{config.EMPLOYEES_TABLE}] WHERE {where}", [key]
        )
        found = cur.fetchone()
        if found is None:
            return None
        matched_name = _clean(found[0])
        if db_values:
            set_sql = ", ".join(f"[{c}] = ?" for c in db_values)
            cur.execute(
                f"UPDATE [{config.EMPLOYEES_TABLE}] SET {set_sql} WHERE {where}",
                list(db_values.values()) + [key],
            )
        conn.commit()
    finally:
        conn.close()
    _clear_employee_caches()
    return employee_index().get(matched_name)


# ---------------------------------------------------------------------------
# CSV backend — writes (append/rewrite the dump files).
# ---------------------------------------------------------------------------
def _append_csv_row(path, values: dict[str, str]) -> None:
    """Append one row to a CSV, matching its existing header order."""
    with path.open(encoding="utf-8-sig", newline="") as fh:
        fieldnames = csv.DictReader(fh).fieldnames or []

    needs_nl = False
    with path.open("rb") as fh:
        fh.seek(0, 2)
        if fh.tell() > 0:
            fh.seek(-1, 2)
            needs_nl = fh.read(1) not in (b"\n", b"\r")

    row = {name: "" for name in fieldnames}
    for col, val in values.items():
        if col in row:
            row[col] = _clean(val)

    with path.open("a", encoding="utf-8", newline="") as fh:
        if needs_nl:
            fh.write("\r\n")
        csv.DictWriter(fh, fieldnames=fieldnames, extrasaction="ignore").writerow(row)


def _csv_append_project(values: dict[str, str]) -> Project:
    _append_csv_row(config.FUNNEL_FILE, values)
    _clear_project_caches()
    return project_index()[values[PROJECT_COLUMNS["project_id"]]]


def _csv_append_projects(rows: list[dict[str, str]]) -> int:
    path = config.FUNNEL_FILE
    with path.open(encoding="utf-8-sig", newline="") as fh:
        fieldnames = csv.DictReader(fh).fieldnames or []

    needs_nl = False
    with path.open("rb") as fh:
        fh.seek(0, 2)
        if fh.tell() > 0:
            fh.seek(-1, 2)
            needs_nl = fh.read(1) not in (b"\n", b"\r")

    with path.open("a", encoding="utf-8", newline="") as fh:
        if needs_nl:
            fh.write("\r\n")
        writer = csv.DictWriter(fh, fieldnames=fieldnames, extrasaction="ignore")
        for values in rows:
            base = {name: "" for name in fieldnames}
            for col, val in values.items():
                if col in base:
                    base[col] = _clean(val)
            writer.writerow(base)

    _clear_project_caches()
    return len(rows)


def _csv_funnel_row(project_id: str) -> dict[str, str] | None:
    with config.FUNNEL_FILE.open(encoding="utf-8-sig", newline="") as fh:
        for row in csv.DictReader(fh):
            if _clean(row.get(PROJECT_COLUMNS["project_id"])) == project_id:
                return row
    return None


def _csv_update_project(project_id: str, updates: dict[str, str]) -> Project | None:
    path = config.FUNNEL_FILE
    with path.open(encoding="utf-8-sig", newline="") as fh:
        reader = csv.DictReader(fh)
        fieldnames = reader.fieldnames or []
        rows = list(reader)

    found = False
    for r in rows:
        if _clean(r.get(PROJECT_COLUMNS["project_id"])) == project_id:
            for col, val in updates.items():
                if col in fieldnames:
                    r[col] = _clean(val)
            found = True
    if not found:
        return None

    with path.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)

    _clear_project_caches()
    return project_index().get(project_id)


def _employee_matches(row: dict, key: str) -> bool:
    """Match a headcount row by email (preferred, unique) or by name."""
    if "@" in key:
        return _clean(row.get(EMPLOYEE_COLUMNS["email"])).lower() == key.lower()
    return _clean(row.get(EMPLOYEE_COLUMNS["name"])) == key


def _csv_append_employee(values: dict[str, str]) -> Employee:
    _append_csv_row(config.HEADCOUNT_FILE, values)
    _clear_employee_caches()
    return employee_index()[values[EMPLOYEE_COLUMNS["name"]]]


def _csv_employee_row(key: str) -> dict[str, str] | None:
    with config.HEADCOUNT_FILE.open(encoding="utf-8-sig", newline="") as fh:
        for row in csv.DictReader(fh):
            if _employee_matches(row, key):
                return row
    return None


def _csv_update_employee(key: str, updates: dict[str, str]) -> Employee | None:
    path = config.HEADCOUNT_FILE
    with path.open(encoding="utf-8-sig", newline="") as fh:
        reader = csv.DictReader(fh)
        fieldnames = reader.fieldnames or []
        rows = list(reader)

    matched_name = ""
    for r in rows:
        if _employee_matches(r, key):
            for col, val in updates.items():
                if col in fieldnames:
                    r[col] = _clean(val)
            matched_name = _clean(r.get(EMPLOYEE_COLUMNS["name"]))
            break
    if not matched_name:
        return None

    with path.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)

    _clear_employee_caches()
    return employee_index().get(matched_name)


# ---------------------------------------------------------------------------
# CSV exports (used by the /export endpoints).
# ---------------------------------------------------------------------------
def _sql_export_csv(table: str, column_map: dict[str, str]) -> str:
    import io

    cols = list(column_map.keys())
    select = ", ".join(f"[{c}]" for c in cols)
    conn = db.get_connection()
    try:
        cur = conn.cursor()
        cur.execute(f"SELECT {select} FROM [{table}]")
        rows = cur.fetchall()
    finally:
        conn.close()

    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow([column_map[c] for c in cols])
    for row in rows:
        writer.writerow(["" if v is None else v for v in row])
    return buf.getvalue()


def export_projects_csv() -> str:
    if active_backend() == "csv":
        return config.FUNNEL_FILE.read_text(encoding="utf-8-sig")
    try:
        return _sql_export_csv(config.PROJECTS_TABLE, PROJECT_COLUMNS)
    except Exception:
        _use_csv()
        return config.FUNNEL_FILE.read_text(encoding="utf-8-sig")


def export_employees_csv() -> str:
    if active_backend() == "csv":
        return config.HEADCOUNT_FILE.read_text(encoding="utf-8-sig")
    try:
        return _sql_export_csv(config.EMPLOYEES_TABLE, EMPLOYEE_COLUMNS)
    except Exception:
        _use_csv()
        return config.HEADCOUNT_FILE.read_text(encoding="utf-8-sig")
