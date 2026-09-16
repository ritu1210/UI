"""JSON API endpoints consumed by the front-end pages."""
from __future__ import annotations

import csv
import io
import re
from collections import Counter, defaultdict

from fastapi import APIRouter, File, HTTPException, UploadFile
from fastapi.responses import Response
from pydantic import BaseModel

from . import config, data, db, store


def _pretty_manager(email: str) -> str:
    """Turn a manager email (tim.bokelman@philips.com) into a readable name."""
    if not email:
        return email
    local = email.split("@")[0]
    parts = [p for p in re.split(r"[._]", local) if p]
    return " ".join(p.capitalize() for p in parts) or email

router = APIRouter(prefix="/api", tags=["api"])


@router.get("/health/db")
def health_db():
    """Report Azure SQL connectivity for the configured _dev tables."""
    result = db.health_check()
    if result.get("status") != "ok":
        raise HTTPException(status_code=503, detail=result)
    return result


class AllocationIn(BaseModel):
    employee: str
    project_id: str
    month: str
    allocation: float


class AllocationPatch(BaseModel):
    employee: str | None = None
    project_id: str | None = None
    month: str | None = None
    allocation: float | None = None


def _serialize(alloc: store.Allocation) -> dict:
    return {
        "id": alloc.id,
        "employee": alloc.employee,
        "project_id": alloc.project_id,
        "project_title": alloc.project_title,
        "month": alloc.month,
        "allocation": alloc.allocation,
    }


@router.get("/employees")
def get_employees(q: str | None = None):
    employees = data.load_employees()
    if q:
        needle = q.lower()
        employees = [e for e in employees if needle in e.name.lower()]
    return [
        {
            "name": e.name,
            "job_title": e.job_title,
            "director": e.director,
            "reporting_manager": e.reporting_manager,
            "location": e.location,
            "country": e.country,
        }
        for e in employees
    ]


@router.get("/projects")
def get_projects(q: str | None = None, limit: int = 50):
    projects = data.load_projects()
    if q:
        needle = q.lower()
        projects = [
            p for p in projects
            if needle in p.project_id.lower() or needle in p.title.lower()
        ]
    return [
        {
            "project_id": p.project_id,
            "title": p.title,
            "bu": p.bu,
            "cluster": p.cluster,
            "project_type": p.project_type,
            "spoc": p.spoc,
            "current_il": p.current_il,
            "is_active": p.is_active or "Yes",
        }
        for p in projects[:limit]
    ]


@router.get("/projects/page")
def get_projects_page(q: str | None = None, active: str | None = None, bu: str | None = None, spoc: str | None = None, page: int = 1, size: int = 25):
    """Paginated project list for the SharePoint-style Manage Projects table."""
    projects = data.load_projects()
    if q:
        needle = q.lower()
        projects = [
            p for p in projects
            if needle in p.project_id.lower() or needle in p.title.lower()
        ]
    if active in ("yes", "no"):
        want_active = active == "yes"
        projects = [p for p in projects if ((p.is_active or "Yes").lower() != "no") == want_active]
    if bu:
        needle_bu = bu.lower()
        projects = [p for p in projects if p.bu.lower() == needle_bu]
    if spoc:
        projects = [p for p in projects if p.spoc == spoc]

    total = len(projects)
    size = max(1, min(size, 200))
    page = max(1, page)
    start = (page - 1) * size
    window = projects[start:start + size]
    return {
        "total": total,
        "page": page,
        "size": size,
        "items": [
            {
                "project_id": p.project_id,
                "title": p.title,
                "bu": p.bu,
                "cluster": p.cluster,
                "project_type": p.project_type,
                "current_il": p.current_il,
                "is_active": p.is_active or "Yes",
                "spoc": p.spoc,
                "funnel_total": p.funnel_total,
                "actual_total": p.actual_total,
            }
            for p in window
        ],
    }


@router.get("/projects/summary")
def get_projects_summary(q: str | None = None, active: str | None = None, bu: str | None = None, spoc: str | None = None):
    """Totals and savings for the current filter set, plus per-BU / per-SPOC counts for the filters."""
    projects = data.load_projects()

    # Dropdown counts respect search + status (but not the BU/SPOC selection) so the filters stay useful.
    scoped = projects
    if q:
        needle = q.lower()
        scoped = [p for p in scoped if needle in p.project_id.lower() or needle in p.title.lower()]
    if active in ("yes", "no"):
        want_active = active == "yes"
        scoped = [p for p in scoped if ((p.is_active or "Yes").lower() != "no") == want_active]

    bu_variants: dict[str, Counter] = defaultdict(Counter)
    spoc_counts: dict[str, int] = defaultdict(int)
    for p in scoped:
        if p.bu:
            bu_variants[p.bu.lower()][p.bu] += 1
        if p.spoc:
            spoc_counts[p.spoc] += 1
    # Merge BU casing variants (DXR / DxR) and drop one-off junk entries.
    bus = []
    for variants in bu_variants.values():
        total = sum(variants.values())
        if total <= 1:
            continue
        bus.append({"bu": variants.most_common(1)[0][0], "count": total})
    bus.sort(key=lambda x: x["bu"].lower())
    spocs = sorted(({"spoc": name, "count": c} for name, c in spoc_counts.items()), key=lambda x: x["spoc"].lower())

    filtered = scoped
    if bu:
        needle_bu = bu.lower()
        filtered = [p for p in filtered if p.bu.lower() == needle_bu]
    if spoc:
        filtered = [p for p in filtered if p.spoc == spoc]
    active_count = sum(1 for p in filtered if (p.is_active or "Yes").lower() != "no")
    return {
        "total": len(filtered),
        "active": active_count,
        "inactive": len(filtered) - active_count,
        "funnel_total": sum(p.funnel_total for p in filtered),
        "actual_total": sum(p.actual_total for p in filtered),
        "bus": bus,
        "spocs": spocs,
    }


@router.get("/managers")
def get_managers():
    counts: dict[str, int] = {}
    for e in data.load_employees():
        if e.reporting_manager:
            key = e.reporting_manager.lower()
            counts[key] = counts.get(key, 0) + 1
    result = [
        {"value": m, "label": _pretty_manager(m), "reportees": c}
        for m, c in counts.items()
    ]
    result.sort(key=lambda r: r["label"].lower())
    return result


@router.get("/allocations")
def get_allocations(manager: str | None = None):
    allocs = store.list_all()
    if manager is not None:
        key = manager.lower()
        emp_index = data.employee_index()
        allocs = [
            a for a in allocs
            if (emp_index.get(a.employee) and emp_index[a.employee].reporting_manager.lower() == key)
        ]
    return [_serialize(a) for a in allocs]


@router.get("/reportees")
def get_reportees(manager: str):
    key = manager.lower()
    return [
        {"name": e.name, "job_title": e.job_title, "location": e.location}
        for e in data.load_employees()
        if e.reporting_manager.lower() == key
    ]


@router.post("/allocations", status_code=201)
def create_allocation(payload: AllocationIn):
    project = data.project_index().get(payload.project_id)
    if not project:
        raise HTTPException(status_code=404, detail="Unknown project id")
    if payload.month not in store.allocatable_months():
        raise HTTPException(status_code=422, detail="Allocations are only allowed for the current or future months")
    alloc = store.create(
        employee=payload.employee,
        project_id=payload.project_id,
        project_title=project.title,
        month=payload.month,
        allocation=payload.allocation,
    )
    return _serialize(alloc)


@router.patch("/allocations/{alloc_id}")
def patch_allocation(alloc_id: int, payload: AllocationPatch):
    fields = payload.model_dump(exclude_none=True)
    if "month" in fields and fields["month"] not in store.allocatable_months():
        raise HTTPException(status_code=422, detail="Allocations are only allowed for the current or future months")
    if "project_id" in fields:
        project = data.project_index().get(fields["project_id"])
        if not project:
            raise HTTPException(status_code=404, detail="Unknown project id")
        fields["project_title"] = project.title
    alloc = store.update(alloc_id, **fields)
    if not alloc:
        raise HTTPException(status_code=404, detail="Allocation not found")
    return _serialize(alloc)


@router.delete("/allocations/{alloc_id}", status_code=204)
def delete_allocation(alloc_id: int):
    if not store.delete(alloc_id):
        raise HTTPException(status_code=404, detail="Allocation not found")


@router.get("/dashboard/user")
def dashboard_user(employee: str | None = None):
    allocations = store.list_all()
    if employee:
        allocations = [a for a in allocations if a.employee == employee]
    by_month: dict[str, float] = defaultdict(float)
    rows = []
    for a in allocations:
        by_month[a.month] += a.allocation
        rows.append(_serialize(a))
    return {
        "rows": rows,
        "by_month": [{"month": m, "total": round(t, 1)} for m, t in sorted(by_month.items())],
        "total_allocations": len(rows),
    }


@router.get("/dashboard/user-breakdown")
def dashboard_user_breakdown(
    manager: str | None = None,
    employee: str | None = None,
    month: str | None = None,
):
    proj_index = data.project_index()
    emp_index = data.employee_index()
    allocs = store.list_all()
    if manager:
        key = manager.lower()
        allocs = [
            a for a in allocs
            if emp_index.get(a.employee) and emp_index[a.employee].reporting_manager.lower() == key
        ]
    if employee:
        allocs = [a for a in allocs if a.employee == employee]
    if month:
        allocs = [a for a in allocs if a.month == month]

    def breakdown(attr: str) -> list[dict]:
        sums: dict[str, float] = defaultdict(float)
        counts: dict[str, int] = defaultdict(int)
        for a in allocs:
            project = proj_index.get(a.project_id)
            label = (getattr(project, attr) if project else "") or "Unassigned"
            sums[label] += a.allocation
            counts[label] += 1
        rows = [{"label": k, "value": round(sums[k] / counts[k], 1)} for k in sums]
        rows.sort(key=lambda r: r["value"], reverse=True)
        return rows[:12]

    return {
        "by_cluster": breakdown("cluster"),
        "by_bu": breakdown("bu"),
        "by_project_type": breakdown("project_type"),
        "by_il": breakdown("current_il"),
        "meta": {
            "allocations": len(allocs),
            "employees": len({a.employee for a in allocs}),
            "avg": round(sum(a.allocation for a in allocs) / len(allocs), 1) if allocs else 0,
        },
    }


# Utilization bands (per reportee, % of their monthly capacity that is booked).
_BENCH_MAX = 0.0       # exactly 0% booked -> on the bench
_UNDER_MAX = 79.0      # 1-79% -> under-utilized (spare capacity)
_HEALTHY_MAX = 100.0   # 80-100% -> healthy
                       # >100% -> over-allocated


def _band(util: float) -> str:
    if util <= _BENCH_MAX:
        return "bench"
    if util <= _UNDER_MAX:
        return "under"
    if util <= _HEALTHY_MAX:
        return "healthy"
    return "over"


@router.get("/dashboard/people-leader")
def dashboard_people_leader(month: str | None = None, manager: str | None = None):
    """Leadership view: team utilization, capacity coverage and bench risk.

    Utilization for a reportee = the sum of their project allocations in the
    selected month (can exceed 100% when double-booked across projects).
    When ``manager`` is supplied the view narrows to that manager's reportees.
    """
    employees = data.load_employees()
    sel_month = month if month in store.MONTHS else store.current_month()

    # Total booked % and project count per (employee, month).
    util_by_emp: dict[tuple[str, str], float] = defaultdict(float)
    projects_in_month: dict[str, int] = defaultdict(int)
    allocs_in_month: dict[str, list] = defaultdict(list)
    proj_index = data.project_index()
    for a in store.list_all():
        util_by_emp[(a.employee, a.month)] += a.allocation
        if a.month == sel_month:
            projects_in_month[a.employee] += 1
            p = proj_index.get(a.project_id)
            allocs_in_month[a.employee].append({
                "project_id": a.project_id,
                "project_title": a.project_title,
                "allocation": a.allocation,
                "current_il": p.current_il if p else "",
                "bu": p.bu if p else "",
                "project_type": p.project_type if p else "",
            })

    # Reportees grouped under their reporting manager (headcount from HC list).
    teams: dict[str, list] = defaultdict(list)
    for e in employees:
        if e.reporting_manager:
            teams[e.reporting_manager].append(e)

    bands = {"bench": 0, "under": 0, "healthy": 0, "over": 0}

    # ---- Single-team view -------------------------------------------------
    if manager:
        key = manager.lower()
        members = [e for e in employees if e.reporting_manager and e.reporting_manager.lower() == key]
        by_reportee = []
        for e in members:
            u = util_by_emp.get((e.name, sel_month), 0.0)
            bands[_band(u)] += 1
            by_reportee.append({
                "name": e.name,
                "job_title": e.job_title,
                "util": round(u, 1),
                "projects": projects_in_month.get(e.name, 0),
                "status": _band(u),
                "employment_status": e.status,
                "allocations": sorted(allocs_in_month.get(e.name, []), key=lambda x: x["allocation"], reverse=True),
            })
        by_reportee.sort(key=lambda r: (r["util"], r["name"].lower()))
        total = len(members)
        total_util = sum(r["util"] for r in by_reportee)
        trend = [
            {
                "month": m,
                "avg_util": round(sum(util_by_emp.get((e.name, m), 0.0) for e in members) / total, 1) if total else 0.0,
            }
            for m in store.MONTHS
        ]
        return {
            "scope": "team",
            "scope_label": _pretty_manager(manager),
            "manager": manager,
            "month": sel_month,
            "months": store.MONTHS,
            "kpis": {
                "reportees": total,
                "avg_util": round(total_util / total, 1) if total else 0,
                "fully_allocated": bands["healthy"] + bands["over"],
                "over_allocated": bands["over"],
                "bench": bands["bench"],
                "sum_allocations": round(total_util),
            },
            "utilization": bands,
            "by_reportee": by_reportee,
            "by_manager": [],
            "trend": trend,
        }

    # ---- Org-wide view ----------------------------------------------------
    total_reportees = 0
    total_util = 0.0
    by_manager = []
    for mgr, members in teams.items():
        utils = [util_by_emp.get((e.name, sel_month), 0.0) for e in members]
        allocated = sum(1 for u in utils if u > 0)
        team_total = sum(utils)
        avg_util = round(team_total / len(members), 1) if members else 0.0
        for u in utils:
            bands[_band(u)] += 1
        total_reportees += len(members)
        total_util += team_total
        by_manager.append({
            "manager": _pretty_manager(mgr),
            "team_size": len(members),
            "allocated": allocated,
            "avg_util": avg_util,
            "total_util": round(team_total, 1),
            "over": sum(1 for u in utils if u > _HEALTHY_MAX),
        })

    by_manager.sort(key=lambda r: r["avg_util"], reverse=True)

    # 12-month utilization trend across every reportee (org-wide average).
    all_members = [e for members in teams.values() for e in members]
    trend = []
    for m in store.MONTHS:
        if all_members:
            avg = sum(util_by_emp.get((e.name, m), 0.0) for e in all_members) / len(all_members)
        else:
            avg = 0.0
        trend.append({"month": m, "avg_util": round(avg, 1)})

    return {
        "scope": "org",
        "scope_label": "All teams",
        "manager": None,
        "month": sel_month,
        "months": store.MONTHS,
        "kpis": {
            "leaders": len(teams),
            "reportees": total_reportees,
            "avg_util": round(total_util / total_reportees, 1) if total_reportees else 0,
            "fully_allocated": bands["healthy"] + bands["over"],
            "over_allocated": bands["over"],
            "bench": bands["bench"],
            "sum_allocations": round(total_util),
        },
        "utilization": bands,
        "by_reportee": [],
        "by_manager": by_manager,
        "trend": trend,
    }


@router.get("/dashboard/reportees-by-status")
def dashboard_reportees_by_status(status: str, month: str | None = None, manager: str | None = None):
    """List the individual reportees in a utilization band (and who they report to).

    ``status`` is one of bench / under / healthy / over, or the group ``fully``
    (healthy + over). Optionally narrowed to a single ``manager``'s team.
    """
    employees = data.load_employees()
    sel_month = month if month in store.MONTHS else store.current_month()

    util_by_emp: dict[str, float] = defaultdict(float)
    projects_in_month: dict[str, int] = defaultdict(int)
    for a in store.list_all():
        if a.month == sel_month:
            util_by_emp[a.employee] += a.allocation
            projects_in_month[a.employee] += 1

    wanted = {"healthy", "over"} if status == "fully" else {status}
    key = manager.lower() if manager else None

    rows = []
    for e in employees:
        if not e.reporting_manager:
            continue
        if key and e.reporting_manager.lower() != key:
            continue
        util = util_by_emp.get(e.name, 0.0)
        band = _band(util)
        if band not in wanted:
            continue
        rows.append({
            "name": e.name,
            "job_title": e.job_title,
            "reporting_manager": _pretty_manager(e.reporting_manager),
            "reporting_manager_email": e.reporting_manager,
            "location": e.location,
            "util": round(util, 1),
            "projects": projects_in_month.get(e.name, 0),
            "status": band,
        })
    rows.sort(key=lambda r: (r["reporting_manager"].lower(), r["name"].lower()))
    return {"month": sel_month, "status": status, "count": len(rows), "rows": rows}


# ---------------------------------------------------------------------------
# PM data entry: append rows to the source CSV files (Funnel / Headcount).
# ---------------------------------------------------------------------------

FUNNEL_COLUMNS = {
    "project_id": "SMRS / Project ID",
    "title": "STET Funnel Project Title",
    "project_type": "Project Type (PRODUCTIVITY, AOS, LCM, NPI, CART PDC, TEST ENG, CONQ, IGM)",
    "bu": "BU",
    "cluster": "Cluster",
    "commodity": "Commodity",
    "current_il": "Current IL",
    "il5_date": "IL5 Date",
    "is_active": "Is Active",
    "parts_dual_sourced": "%23 Parts of Dual Sourced",
    "director": "Director",
    "spoc": "STET SPOC",
    "program_manager": "Program Manager",
    "aos_impact": "AOS Impact  (Enter €0 if not AOS Project)",
    "qn_reduction": "QN Reduction Impact (Enter €0 if not CONQ Project)",
    "procurement_type": "Procurement Type (PROCUREMENT, TCO, N/A)",
    "savings_type": "Savings Type (TCO, CONCEPT, SOURCING, NEGO, NON_12NC, NPP, CONQ, FCP & PPV)",
    "funnel_2025": "2025 Funnel (Euro)",
    "actual_2025": "2025 Actual (Euro)",
    "funnel_2026": "2026 Funnel (Euro)",
    "actual_2026": "2026 Actual (Euro)",
    "funnel_2027": "2027 Funnel (Euro)",
    "actual_2027": "2027 Actual (Euro)",
    "funnel_2028": "2028 Funnel (Euro)",
    "actual_2028": "2028 Actual (Euro)",
    "impacted_parts": "Impacted Parts Added? (YES_NO_N/A) Needs to be YES or N/A for all IL5 Projects",
    "sqe_resources": "Are Any STET SQE Resources Applied to Project? (YES_NO_N/A)",
    "week": "Week",
    "comments": "Comments / Challenges",
}

HEADCOUNT_COLUMNS = {
    "name": "Employee Name (HC)",
    "email": "Email ID",
    "job_title": "Job Title (HC)",
    "job_grade": "Job Grade (HC)",
    "director": "Director",
    "country": "Location Country (HC)",
    "location": "Job Location (HC)",
    "gender": "Diversity (HC)",
    "status": "Employment Status (HC)",
    "start_date": "Employee Start Date (M/D/YYYY)",
    "reporting_manager": "Reporting Manager (HC)",
    "employment_type": "Employment Type (HC)",
}


class FunnelIn(BaseModel):
    project_id: str
    title: str
    project_type: str = ""
    bu: str = ""
    cluster: str = ""
    commodity: str = ""
    current_il: str = ""
    il5_date: str = ""
    is_active: str = "Yes"
    parts_dual_sourced: str = ""
    director: str = ""
    spoc: str = ""
    program_manager: str = ""
    aos_impact: str = ""
    qn_reduction: str = ""
    procurement_type: str = ""
    savings_type: str = ""
    funnel_2025: str = ""
    actual_2025: str = ""
    funnel_2026: str = ""
    actual_2026: str = ""
    funnel_2027: str = ""
    actual_2027: str = ""
    funnel_2028: str = ""
    actual_2028: str = ""
    impacted_parts: str = ""
    sqe_resources: str = ""
    week: str = ""
    comments: str = ""


class HeadcountIn(BaseModel):
    name: str
    email: str
    job_title: str = ""
    job_grade: str = ""
    director: str = ""
    country: str = ""
    location: str = ""
    gender: str = ""
    status: str = ""
    start_date: str = ""
    reporting_manager: str = ""
    employment_type: str = ""


@router.get("/stats")
def stats():
    """High-level counts for the welcome page."""
    projects = data.load_projects()
    employees = data.load_employees()
    managers = {e.reporting_manager.lower() for e in employees if e.reporting_manager}
    bus = {p.bu for p in projects if p.bu}
    return {
        "projects": len(projects),
        "employees": len(employees),
        "leaders": len(managers),
        "business_units": len(bus),
    }


# Values to drop from the data-entry suggestion lists (case-insensitive).
META_BLOCK = {
    "bus": {"all", "events", "funnel activities", "long leave", "l&a", "trainings"},
    "clusters": {"funnel activities", "events", "long leave", "trainings", "l&a"},
    "project_types": {"0"},
    "commodities": {"all", "long leave", "n/a", "key component & electronics"},
    "current_ils": {"ilr"},
    "directors": {"vanish", "nishant", "dan", "nishant dan"},
}


@router.get("/meta")
def meta():
    """Distinct, de-duplicated values used to power the data-entry form suggestions."""
    projects = data.load_projects()
    employees = data.load_employees()

    def clean(values, key=""):
        block = META_BLOCK.get(key, set())
        seen, out = set(), []
        for v in sorted({x for x in values if x}):
            norm = v.strip().lower()
            if norm in block or norm in seen:
                continue
            seen.add(norm)
            out.append(v)
        return out

    return {
        "bus": clean((p.bu for p in projects), "bus"),
        "clusters": clean((p.cluster for p in projects), "clusters"),
        "project_types": clean((p.project_type for p in projects), "project_types"),
        "commodities": clean((p.commodity for p in projects), "commodities"),
        "current_ils": clean((p.current_il for p in projects), "current_ils"),
        "spocs": clean(p.spoc for p in projects),
        "program_managers": clean(p.program_manager for p in projects),
        "directors": clean([p.director for p in projects] + [e.director for e in employees], "directors"),
        "countries": clean(e.country for e in employees),
        "locations": clean(e.location for e in employees),
        "job_titles": clean(e.job_title for e in employees),
        "job_grades": clean(e.job_grade for e in employees),
        "employment_types": clean(e.employment_type for e in employees),
        "statuses": clean(e.status for e in employees),
        "reporting_managers": clean(e.reporting_manager for e in employees),
    }


@router.post("/funnel", status_code=201)
def add_funnel_project(payload: FunnelIn):
    if not payload.project_id.strip() or not payload.title.strip():
        raise HTTPException(status_code=422, detail="Project ID and Title are required")
    if data.project_index().get(payload.project_id.strip()):
        raise HTTPException(status_code=409, detail="A project with this ID already exists")
    values = {FUNNEL_COLUMNS[k]: v for k, v in payload.model_dump().items()}
    project = data.append_project(values)
    return {
        "project_id": project.project_id,
        "title": project.title,
        "bu": project.bu,
        "cluster": project.cluster,
        "project_type": project.project_type,
        "spoc": project.spoc,
    }


# ---------------------------------------------------------------------------
# Bulk import: add many funnel projects from a spreadsheet (CSV / Excel /
# Smartsheet export). Column headers are matched flexibly by name.
# ---------------------------------------------------------------------------

# Extra human-friendly header names accepted for spreadsheet / Smartsheet imports.
# (The exact CSV column names and the API field keys are always matched too.)
FUNNEL_HEADER_ALIASES = {
    "project_id": ["project id", "smrs", "smrs id", "smrs / project id", "smrs/project id", "id"],
    "title": ["project title", "title", "name", "project name", "funnel project title"],
    "project_type": ["project type", "type"],
    "bu": ["business unit"],
    "cluster": ["cluster"],
    "commodity": ["commodity"],
    "current_il": ["il", "current integration level", "integration level"],
    "il5_date": ["il5 date"],
    "is_active": ["active", "status"],
    "director": ["director"],
    "spoc": ["spoc", "stet spoc"],
    "program_manager": ["program manager", "pm"],
    "comments": ["comments", "challenges", "comments / challenges"],
}

_MAX_IMPORT_BYTES = 12 * 1024 * 1024  # 12 MB


def _norm_header(value: str | None) -> str:
    """Normalise a spreadsheet header for tolerant matching."""
    return " ".join(str(value or "").split()).strip().lower().rstrip(":*").strip()


def _funnel_header_map() -> dict[str, str]:
    """Map every recognised (normalised) header to its funnel field key."""
    mapping: dict[str, str] = {}
    for key, col in FUNNEL_COLUMNS.items():
        mapping[_norm_header(col)] = key
        mapping[_norm_header(key)] = key
    for key, aliases in FUNNEL_HEADER_ALIASES.items():
        for alias in aliases:
            mapping[_norm_header(alias)] = key
    return mapping


def _read_upload_rows(filename: str, raw: bytes) -> list[dict[str, str]]:
    """Parse an uploaded CSV or XLSX file into a list of header->value dicts."""
    name = (filename or "").lower()
    if name.endswith((".csv", ".txt")):
        text = raw.decode("utf-8-sig", errors="replace")
        return [dict(r) for r in csv.DictReader(io.StringIO(text))]
    if name.endswith((".xlsx", ".xlsm")):
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise HTTPException(status_code=500, detail="Excel support is not installed on the server (pip install openpyxl)")
        try:
            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        except Exception:
            raise HTTPException(status_code=422, detail="Could not read the Excel file — is it a valid .xlsx?")
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        try:
            header = next(rows_iter)
        except StopIteration:
            wb.close()
            return []
        headers = [str(h).strip() if h is not None else "" for h in header]
        out: list[dict[str, str]] = []
        for values in rows_iter:
            row: dict[str, str] = {}
            for i, head in enumerate(headers):
                if not head:
                    continue
                cell = values[i] if i < len(values) else None
                row[head] = "" if cell is None else str(cell)
            if any(v.strip() for v in row.values()):
                out.append(row)
        wb.close()
        return out
    raise HTTPException(status_code=422, detail="Unsupported file type. Upload a .csv or .xlsx file")


@router.post("/funnel/import")
async def import_funnel_projects(file: UploadFile = File(...)):
    """Bulk-add funnel projects from an uploaded spreadsheet.

    Rows without a Project ID and Title are skipped, as are IDs that already
    exist or repeat within the file. Returns a per-category summary.
    """
    raw = await file.read()
    if not raw:
        raise HTTPException(status_code=422, detail="The uploaded file is empty")
    if len(raw) > _MAX_IMPORT_BYTES:
        raise HTTPException(status_code=413, detail="File is too large (max 12 MB)")

    rows = _read_upload_rows(file.filename or "", raw)
    if not rows:
        raise HTTPException(status_code=422, detail="No data rows were found in the file")

    header_map = _funnel_header_map()
    existing = set(data.project_index().keys())
    seen_ids: set[str] = set()
    to_add: list[dict[str, str]] = []
    skipped_existing = duplicates = invalid = 0
    errors: list[str] = []

    for line, raw_row in enumerate(rows, start=2):  # row 1 is the header
        mapped: dict[str, str] = {}
        for header, val in raw_row.items():
            key = header_map.get(_norm_header(header))
            if key:
                mapped[key] = "" if val is None else str(val).strip()
        pid = (mapped.get("project_id") or "").strip()
        title = (mapped.get("title") or "").strip()
        if not pid or not title:
            invalid += 1
            if len(errors) < 15:
                errors.append(f"Row {line}: missing Project ID or Title")
            continue
        if pid in existing:
            skipped_existing += 1
            continue
        if pid in seen_ids:
            duplicates += 1
            continue
        seen_ids.add(pid)
        to_add.append({FUNNEL_COLUMNS[k]: v for k, v in mapped.items() if k in FUNNEL_COLUMNS})

    added = data.append_projects(to_add)
    return {
        "added": added,
        "skipped_existing": skipped_existing,
        "duplicates": duplicates,
        "invalid": invalid,
        "total_rows": len(rows),
        "errors": errors,
    }


class FunnelUpdate(BaseModel):
    project_id: str
    title: str | None = None
    project_type: str | None = None
    bu: str | None = None
    cluster: str | None = None
    commodity: str | None = None
    current_il: str | None = None
    il5_date: str | None = None
    is_active: str | None = None
    parts_dual_sourced: str | None = None
    director: str | None = None
    spoc: str | None = None
    program_manager: str | None = None
    aos_impact: str | None = None
    qn_reduction: str | None = None
    procurement_type: str | None = None
    savings_type: str | None = None
    funnel_2025: str | None = None
    actual_2025: str | None = None
    funnel_2026: str | None = None
    actual_2026: str | None = None
    funnel_2027: str | None = None
    actual_2027: str | None = None
    funnel_2028: str | None = None
    actual_2028: str | None = None
    impacted_parts: str | None = None
    sqe_resources: str | None = None
    week: str | None = None
    comments: str | None = None


@router.get("/project")
def get_project_detail(project_id: str):
    """Full editable field set for a single funnel project (SharePoint-style edit)."""
    row = data.funnel_row(project_id)
    if row is None:
        raise HTTPException(status_code=404, detail="Project not found")
    fields = {k: " ".join((row.get(col) or "").split()).strip() for k, col in FUNNEL_COLUMNS.items()}
    return {"project_id": project_id, "fields": fields}


@router.patch("/funnel")
def update_funnel_project(payload: FunnelUpdate):
    pid = payload.project_id.strip()
    if not data.project_index().get(pid):
        raise HTTPException(status_code=404, detail="Project not found")
    updates: dict[str, str] = {}
    for k, v in payload.model_dump(exclude={"project_id"}, exclude_none=True).items():
        col = FUNNEL_COLUMNS.get(k)
        if col:
            updates[col] = v
    if not updates:
        raise HTTPException(status_code=422, detail="Nothing to update")
    p = data.update_project(pid, updates)
    if p is None:
        raise HTTPException(status_code=404, detail="Project not found")
    return {
        "project_id": p.project_id,
        "title": p.title,
        "bu": p.bu,
        "cluster": p.cluster,
        "project_type": p.project_type,
        "current_il": p.current_il,
        "is_active": p.is_active,
    }


@router.post("/headcount", status_code=201)
def add_headcount(payload: HeadcountIn):
    if not payload.name.strip() or not payload.email.strip():
        raise HTTPException(status_code=422, detail="Employee Name and Email are required")
    email = payload.email.strip().lower()
    if any(e.email.lower() == email for e in data.load_employees() if e.email):
        raise HTTPException(status_code=409, detail="An employee with this email already exists")
    values = {HEADCOUNT_COLUMNS[k]: v for k, v in payload.model_dump().items()}
    employee = data.append_employee(values)
    return {
        "name": employee.name,
        "job_title": employee.job_title,
        "reporting_manager": employee.reporting_manager,
        "location": employee.location,
        "email": employee.email,
    }


@router.get("/funnel/export")
def export_funnel():
    """Download the full funnel data (live from the database) as a CSV file."""
    return Response(
        content=data.export_projects_csv(),
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=STET_Funnel.csv"},
    )


@router.get("/headcount/export")
def export_headcount():
    """Download the full headcount data (live from the database) as a CSV file."""
    return Response(
        content=data.export_employees_csv(),
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=STET_Headcount.csv"},
    )


def _emp_key(e) -> str:
    return e.email or e.name


@router.get("/employees/page")
def get_employees_page(q: str | None = None, director: str | None = None, manager: str | None = None,
                       country: str | None = None, page: int = 1, size: int = 25):
    """Paginated headcount list for the Manage People table."""
    employees = data.load_employees()
    if q:
        needle = q.lower()
        employees = [e for e in employees if needle in e.name.lower() or (e.email and needle in e.email.lower())]
    if director:
        employees = [e for e in employees if e.director == director]
    if manager:
        employees = [e for e in employees if e.reporting_manager == manager]
    if country:
        employees = [e for e in employees if e.country == country]

    total = len(employees)
    size = max(1, min(size, 200))
    page = max(1, page)
    start = (page - 1) * size
    window = employees[start:start + size]
    return {
        "total": total,
        "page": page,
        "size": size,
        "items": [
            {
                "key": _emp_key(e),
                "name": e.name,
                "email": e.email,
                "job_title": e.job_title,
                "job_grade": e.job_grade,
                "director": e.director,
                "reporting_manager": e.reporting_manager,
                "country": e.country,
                "location": e.location,
                "status": e.status,
                "employment_type": e.employment_type,
            }
            for e in window
        ],
    }


# Canonical people-leadership directors (the real four; raw data has junk entries).
PEOPLE_DIRECTORS = ["Otto", "Phil", "Raj", "Tim"]


@router.get("/employees/summary")
def get_employees_summary(q: str | None = None, director: str | None = None, manager: str | None = None,
                          country: str | None = None):
    """Totals and filter option lists for the Manage People page."""
    employees = data.load_employees()

    scoped = employees
    if q:
        needle = q.lower()
        scoped = [e for e in scoped if needle in e.name.lower() or (e.email and needle in e.email.lower())]

    dir_counts: dict[str, int] = defaultdict(int)
    mgr_counts: dict[str, int] = defaultdict(int)
    country_counts: dict[str, int] = defaultdict(int)
    for e in scoped:
        if e.director:
            dir_counts[e.director] += 1
        if e.reporting_manager:
            mgr_counts[e.reporting_manager] += 1
        if e.country:
            country_counts[e.country] += 1
    directors = [{"value": d, "count": dir_counts.get(d, 0)} for d in PEOPLE_DIRECTORS]
    managers = sorted(
        ({"value": k, "label": _pretty_manager(k), "count": c} for k, c in mgr_counts.items()),
        key=lambda x: x["label"].lower(),
    )
    countries = sorted(({"value": k, "count": c} for k, c in country_counts.items()), key=lambda x: x["value"].lower())

    filtered = scoped
    if director:
        filtered = [e for e in filtered if e.director == director]
    if manager:
        filtered = [e for e in filtered if e.reporting_manager == manager]
    if country:
        filtered = [e for e in filtered if e.country == country]
    return {
        "total": len(filtered),
        "managers_count": len({e.reporting_manager for e in filtered if e.reporting_manager}),
        "directors_count": len(PEOPLE_DIRECTORS),
        "countries_count": len({e.country for e in filtered if e.country}),
        "directors": directors,
        "managers": managers,
        "countries": countries,
    }


@router.get("/employee")
def get_employee_detail(key: str):
    """Full editable field set for a single employee (by email or name)."""
    row = data.employee_row(key)
    if row is None:
        raise HTTPException(status_code=404, detail="Employee not found")
    fields = {k: " ".join((row.get(col) or "").split()).strip() for k, col in HEADCOUNT_COLUMNS.items()}
    return {"key": key, "fields": fields}


class HeadcountUpdate(BaseModel):
    key: str
    name: str | None = None
    email: str | None = None
    job_title: str | None = None
    job_grade: str | None = None
    director: str | None = None
    country: str | None = None
    location: str | None = None
    gender: str | None = None
    status: str | None = None
    start_date: str | None = None
    reporting_manager: str | None = None
    employment_type: str | None = None


@router.patch("/headcount")
def update_headcount(payload: HeadcountUpdate):
    key = payload.key.strip()
    if data.employee_row(key) is None:
        raise HTTPException(status_code=404, detail="Employee not found")
    updates: dict[str, str] = {}
    for k, v in payload.model_dump(exclude={"key"}, exclude_none=True).items():
        col = HEADCOUNT_COLUMNS.get(k)
        if col:
            updates[col] = v
    if not updates:
        raise HTTPException(status_code=422, detail="Nothing to update")
    e = data.update_employee(key, updates)
    if e is None:
        raise HTTPException(status_code=404, detail="Employee not found")
    return {
        "key": _emp_key(e),
        "name": e.name,
        "email": e.email,
        "job_title": e.job_title,
        "reporting_manager": e.reporting_manager,
        "location": e.location,
        "status": e.status,
    }


@router.get("/dashboard/business-unit")
def dashboard_business_unit(month: str | None = None):
    """Business-unit view: allocation, headcount and project spread per BU."""
    proj_index = data.project_index()
    sel_month = month if month in store.MONTHS else None  # None => all months

    allocs = store.list_all()
    scoped = [a for a in allocs if a.month == sel_month] if sel_month else allocs

    def bu_of(alloc) -> str:
        project = proj_index.get(alloc.project_id)
        return project.bu if project and project.bu else "Unassigned"

    # BU roster: everyone ever allocated to the BU (denominator for staffing).
    roster: dict[str, set] = defaultdict(set)
    for a in allocs:
        roster[bu_of(a)].add(a.employee)

    summary: dict[str, dict] = {}
    for a in scoped:
        bu = bu_of(a)
        bucket = summary.setdefault(bu, {"bu": bu, "employees": set(), "projects": set(), "total": 0.0, "count": 0})
        bucket["employees"].add(a.employee)
        bucket["projects"].add(a.project_id)
        bucket["total"] += a.allocation
        bucket["count"] += 1

    by_bu = []
    for bucket in summary.values():
        by_bu.append({
            "bu": bucket["bu"],
            "employees": len(bucket["employees"]),
            "headcount": len(roster.get(bucket["bu"], ())),
            "projects": len(bucket["projects"]),
            "allocations": bucket["count"],
            "total_alloc": round(bucket["total"], 1),
            "avg_allocation": round(bucket["total"] / bucket["count"], 1) if bucket["count"] else 0,
        })
    by_bu.sort(key=lambda r: r["employees"], reverse=True)

    # 12-month total allocation trend across all business units.
    trend_by_month: dict[str, float] = defaultdict(float)
    for a in allocs:
        trend_by_month[a.month] += a.allocation
    trend = [{"month": m, "total": round(trend_by_month.get(m, 0.0), 1)} for m in store.MONTHS]

    total_employees = len(data.load_employees())
    allocated_employees = len({a.employee for a in scoped})

    return {
        "month": sel_month,
        "months": store.MONTHS,
        "kpis": {
            "business_units": len(by_bu),
            "total_employees": total_employees,
            "employees": allocated_employees,
            "bench": max(total_employees - allocated_employees, 0),
            "projects": len({a.project_id for a in scoped}),
        },
        "by_bu": by_bu,
        "trend": trend,
    }


@router.get("/kpi/funnel")
def kpi_funnel():
    """Funnel records + distinct slicer options for the Financials/Projects dashboards."""
    records = data.kpi_records()

    def distinct(key: str) -> list[str]:
        return sorted({r[key] for r in records if r.get(key)}, key=lambda s: s.lower())

    filters = {
        k: distinct(k)
        for k in (
            "director", "cluster", "bu", "project_type", "commodity",
            "savings_type", "procurement_type", "current_il", "spoc", "program_manager",
        )
    }
    return {"count": len(records), "records": records, "filters": filters}


@router.get("/kpi/resources")
def kpi_resources():
    """FTE resource-allocation records + slicer options for the Resources dashboard."""
    records = data.resource_records()

    def distinct(key: str) -> list[str]:
        return sorted({r[key] for r in records if r.get(key)}, key=lambda s: s.lower())

    filters = {
        k: distinct(k)
        for k in (
            "proj_others", "director", "cluster", "bu", "region", "project_type",
            "commodity", "savings_type", "fte_type", "status", "reporting_manager",
        )
    }
    total_resource = sum(r["resource_amount"] for r in records)
    return {
        "count": len(records),
        "records": records,
        "filters": filters,
        "total_resource": round(total_resource, 1),
    }
